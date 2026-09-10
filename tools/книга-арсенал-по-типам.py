# -*- coding: utf-8 -*-
"""КНИГА АРСЕНАЛА — листы по типам плюс сводка эффектов.

Карта арсенала состоит из двух половин, и они разной природы:

  * ВЕРХ — разовый утиль: карту сжигают, эффект срабатывает один раз;
  * НИЗ — то, что работает, пока карта УСТАНОВЛЕНА. Он бывает двух видов:
    ПОСТ (пассив, действует сам) и СПЕЦ (спец-действие, которое надо
    разыграть в свой ход, и за него конкурируют все прочие СПЕЦ).

Карты делятся ещё и на обычные и НАЧАЛЬНЫЕ (kind: starting) — последние
раздаются на подготовке и в общую колоду не входят.

Поэтому листов четыре: обычные ПОСТ, обычные СПЕЦ, начальные, и сводка.
Разрез именно такой, потому что решение дизайнера «чего в наборе не хватает»
принимается внутри пары «обычная/начальная × ПОСТ/СПЕЦ», а не по алфавиту.

Сводный лист отвечает на вопрос «какой эффект где стоит»: сколько карт его
несут и НОМЕРА этих карт — отдельно по верхам и по низам. Номера сквозные по
всему набору (1..40) и не сбиваются при переходе с листа на лист.

Версия набора НЕ ЗАШИТА: берётся из свода, который играет движок
(GameConfig.DEFAULT_RULESET -> content_versions.arsenal), см. tools/версии.py.
Захардкоженная версия однажды уже подвела — книга заданий полгода собиралась
из каталога 1.15.0, когда в игре стоял 1.17.0.

Запуск: python tools/книга-арсенал-по-типам.py
"""
import sys
from collections import Counter, defaultdict

sys.path.insert(0, 'tools')
import версии                                    # noqa: E402
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ТОНКАЯ = Side(style='thin', color='FFB0B0B0')
РАМКА = Border(left=ТОНКАЯ, right=ТОНКАЯ, top=ТОНКАЯ, bottom=ТОНКАЯ)


def шапку(ws, заголовки, ширины, цвет):
    ws.append(заголовки)
    for i, (з, ш) in enumerate(zip(заголовки, ширины), start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor=цвет)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = РАМКА
        ws.column_dimensions[get_column_letter(i)].width = ш
    ws.row_dimensions[1].height = 34


def лист(wb, имя, цвет, карты):
    ws = wb.create_sheet(имя[:31])
    заголовки = ['№', 'id', 'НАЗВАНИЕ', 'ВЕРХ: эффект', 'ВЕРХ: что делает',
                 'низ', 'НИЗ: способность', 'НИЗ: что делает', 'описание']
    шапку(ws, заголовки, [4, 8, 22, 20, 44, 7, 24, 50, 58], цвет)
    for строка, (n, к) in enumerate(карты, start=2):
        верх = к.get('top') or {}
        низ = к.get('bottom') or {}
        ws.append([n, к['id'], к.get('name', ''), верх.get('effect', ''),
                   верх.get('label', ''), низ.get('kind', ''),
                   низ.get('passive', низ.get('effect', '')),
                   низ.get('label', ''), к.get('описание', '')])
        for i in range(1, len(заголовки) + 1):
            c = ws.cell(row=строка, column=i)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = РАМКА
        ws.cell(row=строка, column=3).font = Font(bold=True)
    ws.append([])
    ws.append(['', f'карт: {len(карты)}'])
    ws.cell(row=len(карты) + 3, column=2).font = Font(bold=True)
    ws.freeze_panes = 'A2'


def сводка(wb, карты):
    ws = wb.create_sheet('Эффекты — сводка')
    шапку(ws, ['половина', 'эффект / способность', 'карт', 'НОМЕРА КАРТ',
               'что делает (по первой карте)'],
          [10, 32, 6, 30, 70], 'FF3C3160')
    строка = 1
    for половина, ключ in (('ВЕРХ', 'top'), ('НИЗ', 'bottom')):
        счёт = Counter()
        номера = defaultdict(list)
        подпись = {}
        for n, к in карты:
            ч = к.get(ключ) or {}
            имя = ч.get('passive') or ч.get('effect') or '—'
            счёт[имя] += 1
            номера[имя].append(n)
            подпись.setdefault(имя, ч.get('label', ''))
        for имя, сколько in счёт.most_common():
            ws.append([половина, имя, сколько,
                       ', '.join(str(x) for x in номера[имя]), подпись[имя]])
            строка += 1
            for i in range(1, 6):
                c = ws.cell(row=строка, column=i)
                c.alignment = Alignment(vertical='top', wrap_text=True)
                c.border = РАМКА
            ws.cell(row=строка, column=2).font = Font(bold=True)
        ws.append(['', f'разных эффектов: {len(счёт)}', sum(счёт.values()), '', ''])
        строка += 1
        for c in ws[строка]:
            c.font = Font(bold=True)
        ws.append([])
        строка += 1
    ws.freeze_panes = 'A2'


def main():
    свод, в, набор, путь = версии.каталог('arsenal')
    цель = f'docs/АРСЕНАЛ {в} — по типам.xlsx'
    карты = [(n, к) for n, к in enumerate(набор, start=1)]

    def отбор(kind, низKind=None):
        return [(n, к) for n, к in карты
                if к.get('kind') == kind
                and (низKind is None or (к.get('bottom') or {}).get('kind') == низKind)]

    wb = Workbook()
    wb.remove(wb.active)
    лист(wb, 'Обычные — ПОСТ', 'FF2F5D50', отбор('regular', 'POST'))
    лист(wb, 'Обычные — СПЕЦ', 'FF6D4C2F', отбор('regular', 'SPEC'))
    лист(wb, 'Начальные', 'FF2F4F6D', отбор('starting'))
    сводка(wb, карты)

    учтено = sum(len(отбор(*а)) for а in
                 (('regular', 'POST'), ('regular', 'SPEC'), ('starting',)))
    if учтено != len(карты):
        raise SystemExit(f'по листам разошлось: {учтено} из {len(карты)} карт')
    wb.save(цель)
    print(f'готово: {цель}')
    print(f'  свод {свод}, набор {в}, источник {путь}')
    print(f'  карт {len(карты)}, листов {len(wb.sheetnames)}')


if __name__ == '__main__':
    main()
