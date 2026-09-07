# -*- coding: utf-8 -*-
"""КНИГА ЗАДАНИЙ 1.15.0 — лист на каждое семейство плюс сводка эффектов.

Отличается от «ЗАДАНИЯ 1.15.0 — карты целиком» тем, ЗАЧЕМ её открывают. Та —
одна простыня на сорок карт, чтобы читать колоду подряд. Эта — по листу на
семейство, чтобы работать внутри семейства, и сводный лист, который отвечает на
вопрос «какой эффект где стоит»: сколько карт его несут и НОМЕРА этих карт.

Номера сквозные по всей колоде (1..40) и совпадают с номерами в простыне —
иначе двумя таблицами нельзя было бы пользоваться рядом.

Данные берутся из tools/таблица-задания-1-15-0.py: там живёт и колода, и
распределение утилей, и проверка цен. Второй копии списка карт быть не должно.

Запуск: python tools/книга-задания-по-семействам.py
"""
import importlib.util
import io
import sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ЦЕЛЬ = 'docs/ЗАДАНИЯ 1.15.0 — по семействам.xlsx'
ИСТОЧНИК = 'tools/таблица-задания-1-15-0.py'

ТОНКАЯ = Side(style='thin', color='FFB0B0B0')
РАМКА = Border(left=ТОНКАЯ, right=ТОНКАЯ, top=ТОНКАЯ, bottom=ТОНКАЯ)
ШАПКА = 'FF2F5D50'


def загрузить():
    """Взять К и УТИЛИ из таблицы-простыни, не выполняя её main()."""
    spec = importlib.util.spec_from_file_location('таблица1150', ИСТОЧНИК)
    м = importlib.util.module_from_spec(spec)
    sys.modules['таблица1150'] = м
    spec.loader.exec_module(м)
    return м.К, м.УТИЛИ


def шапку(ws, заголовки, ширины, цвет=ШАПКА):
    ws.append(заголовки)
    for i, (з, ш) in enumerate(zip(заголовки, ширины), start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor=цвет)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = РАМКА
        ws.column_dimensions[get_column_letter(i)].width = ш
    ws.row_dimensions[1].height = 34


def main():
    К, УТИЛИ = загрузить()
    if len(К) != 40:
        raise SystemExit(f'заданий {len(К)}, а надо 40')

    # Сквозной номер закрепляется ЗДЕСЬ и дальше только читается.
    пронумерованные = [(n, *карта) for n, карта in enumerate(К, start=1)]

    wb = Workbook()
    wb.remove(wb.active)

    # ---- ЛИСТ НА КАЖДОЕ СЕМЕЙСТВО ----------------------------------------
    семьи = []
    for r in пронумерованные:
        if r[1] not in семьи:
            семьи.append(r[1])
    заголовки = ['№', 'тип', 'НИЗ: условие', 'НИЗ: усиление', 'трофеи', 'монеты',
                 'боеп.', 'зад.', 'цена', 'усиленная награда', 'код', 'тип верха',
                 'ВЕРХ: утиль', 'зачем такая пара']
    ширины = [4, 12, 52, 44, 7, 7, 7, 6, 7, 17, 5, 15, 40, 52]

    for семья in семьи:
        ws = wb.create_sheet(семья[:31])
        шапку(ws, заголовки, ширины)
        строк = 0
        for (n, с, тип, условие, усил, награда, цена, усилНагр, код, зачем) \
                in пронумерованные:
            if с != семья:
                continue
            трофеи, монеты, боеп, зад = награда
            видВерха, действиеВерха, _ = УТИЛИ[код]
            ws.append([n, тип, условие, усил, трофеи or '', монеты or '', боеп or '',
                       зад or '', цена, усилНагр, код, видВерха, действиеВерха, зачем])
            строк += 1
            for i in range(1, len(заголовки) + 1):
                c = ws.cell(row=строк + 1, column=i)
                c.alignment = Alignment(vertical='top', wrap_text=True)
                c.border = РАМКА
            ws.cell(row=строк + 1, column=11).font = Font(bold=True)
        цены = [r[6] for r in пронумерованные if r[1] == семья]
        ws.append([])
        ws.append(['', f'карт: {строк}', f'средняя цена: {sum(цены) / len(цены):.2f}'])
        for c in ws[строк + 3]:
            c.font = Font(bold=True)
        ws.freeze_panes = 'A2'

    # ---- СВОДКА ЭФФЕКТОВ --------------------------------------------------
    ws = wb.create_sheet('Эффекты — сводка')
    шапку(ws, ['код', 'тип верха', 'ЭФФЕКТ (утиль верха)', 'карт',
               'НОМЕРА КАРТ', 'семейства, где стоит'],
          [6, 16, 62, 6, 26, 40], 'FF3C3160')
    счёт = Counter(r[8] for r in пронумерованные)
    номера = defaultdict(list)
    семьиУтиля = defaultdict(list)
    for r in пронумерованные:
        номера[r[8]].append(r[0])
        if r[1] not in семьиУтиля[r[8]]:
            семьиУтиля[r[8]].append(r[1])
    строк = 0
    for код, (вид, действие, надо) in УТИЛИ.items():
        if счёт[код] != надо:
            raise SystemExit(f'утиль {код}: {счёт[код]} карт вместо {надо}')
        ws.append([код, вид, действие, счёт[код],
                   ', '.join(str(x) for x in номера[код]),
                   ', '.join(семьиУтиля[код])])
        строк += 1
        for i in range(1, 7):
            c = ws.cell(row=строк + 1, column=i)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = РАМКА
        ws.cell(row=строк + 1, column=1).font = Font(bold=True)
    ws.append([])
    ws.append(['', 'ВСЕГО', '', sum(счёт.values()), '', ''])
    for c in ws[строк + 3]:
        c.font = Font(bold=True)
    ws.freeze_panes = 'A2'

    # ---- РАЗРЕЗЫ: семейства и типы ----------------------------------------
    ws2 = wb.create_sheet('Разрезы')
    шапку(ws2, ['разрез', 'значение', 'карт', 'средняя цена', 'номера карт'],
          [16, 20, 7, 13, 46], 'FF3C3160')
    строка = 1
    for имя, ключ in (('семейство', 1), ('тип задания', 2)):
        значения = []
        for r in пронумерованные:
            if r[ключ] not in значения:
                значения.append(r[ключ])
        for з in значения:
            свои = [r for r in пронумерованные if r[ключ] == з]
            цены = [r[6] for r in свои]
            ws2.append([имя, з, len(свои), round(sum(цены) / len(цены), 2),
                        ', '.join(str(r[0]) for r in свои)])
            строка += 1
            for i in range(1, 6):
                c = ws2.cell(row=строка, column=i)
                c.alignment = Alignment(vertical='top', wrap_text=True)
                c.border = РАМКА
        ws2.append([])
        строка += 1
    ws2.freeze_panes = 'A2'

    wb.save(ЦЕЛЬ)
    print(f'готово: {ЦЕЛЬ}; листов {len(wb.sheetnames)} '
          f'({len(семьи)} семейств + сводка + разрезы)')


if __name__ == '__main__':
    main()
