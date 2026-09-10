# -*- coding: utf-8 -*-
"""КНИГА СУПЕР-ЗАДАНИЙ — все карты плюс сводка разрезов.

У супер-задания три части, и они разной природы:

  * ТРЕБОВАНИЕ — что надо сделать, чтобы карту развернуть;
  * НАГРАДА — что даётся в момент выполнения, один раз;
  * МНОЖИТЕЛЬ — сколько ПО карта принесёт в финале, и вот он-то и есть
    главное: супер-задание берут ради множителя, а требование только
    открывает к нему доступ.

Листа два, а не семь, как у обычных заданий: поля семейства у супер-заданий в
данных НЕТ, карт всего 17, и выдумывать разрезы за дизайнера незачем. Сводка
режет набор тем, что можно вывести из самих данных, а не на глаз.

Версия каталога НЕ ЗАШИТА: берётся из свода, который играет движок
(GameConfig.DEFAULT_RULESET -> content_versions.super_objectives).

Запуск: python tools/книга-супер-задания.py
"""
import re
import sys
from collections import Counter, defaultdict

sys.path.insert(0, 'tools')
import версии                                    # noqa: E402
from openpyxl import Workbook                    # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402
from openpyxl.utils import get_column_letter     # noqa: E402

ТОНКАЯ = Side(style='thin', color='FFB0B0B0')
РАМКА = Border(left=ТОНКАЯ, right=ТОНКАЯ, top=ТОНКАЯ, bottom=ТОНКАЯ)

# ЧТО СЧИТАЕТ МНОЖИТЕЛЬ. Выводится из текста по ключевым словам, а не
# проставлено вручную: правило разметки видно здесь целиком и проверяемо.
СЧИТАЕТ = [
    ('войска',        r'род(а|ов)? войск|сво[её] войско|своих войск'),
    ('здания',        r'здани'),
    ('трофеи и бой',  r'трофе|уничтоженн'),
    ('карты',         r'задани|арсенал'),
    ('запасы',        r'келемий|монет'),
    ('треки науки',   r'трек'),
    ('порядок хода',  r'первым игроком'),
    ('ЦУ',            r'ЦУ'),
]


def что_считает(текст):
    for имя, шаблон in СЧИТАЕТ:
        if re.search(шаблон, текст, re.I):
            return имя
    return 'прочее'


def потолок(текст):
    m = re.search(r'[Пп]отолок\s+(\d+)', текст)
    return int(m.group(1)) if m else None


def за_единицу(текст):
    m = re.search(r'(\d+)\s*ПО за', текст)
    return int(m.group(1)) if m else None


def шапку(ws, заголовки, ширины, цвет):
    ws.append(заголовки)
    for i, (з, ш) in enumerate(zip(заголовки, ширины), start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor=цвет)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = РАМКА
        ws.column_dimensions[get_column_letter(i)].width = ш
    ws.row_dimensions[1].height = 34


def лист_карт(wb, карты):
    ws = wb.create_sheet('Супер-задания')
    ws.sheet_properties.tabColor = '5A2F6D'
    заголовки = ['№', 'id', 'НАЗВАНИЕ', 'ТРЕБОВАНИЕ', 'НАГРАДА за выполнение',
                 'МНОЖИТЕЛЬ в финале', 'считает', 'ПО за единицу', 'потолок',
                 'описание']
    шапку(ws, заголовки, [4, 8, 24, 48, 48, 48, 14, 8, 8, 58], 'FF5A2F6D')
    for строка, (n, к) in enumerate(карты, start=2):
        м = к.get('multiplier', '')
        ws.append([n, к['id'], к.get('name', ''), к.get('requirement', ''),
                   к.get('reward', ''), м, что_считает(м),
                   за_единицу(м), потолок(м), к.get('описание', '')])
        for i in range(1, len(заголовки) + 1):
            c = ws.cell(row=строка, column=i)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = РАМКА
        ws.cell(row=строка, column=3).font = Font(bold=True)
    ws.append([])
    ws.append(['', f'карт: {len(карты)}'])
    ws.cell(row=len(карты) + 3, column=2).font = Font(bold=True)
    ws.freeze_panes = 'A2'


def сводка(wb, карты):
    ws = wb.create_sheet('Разрезы — сводка')
    ws.sheet_properties.tabColor = '3C3160'
    шапку(ws, ['разрез', 'значение', 'карт', 'НОМЕРА КАРТ'],
          [22, 26, 6, 40], 'FF3C3160')
    строка = 1

    def блок(название, ключ):
        nonlocal строка
        счёт, номера = Counter(), defaultdict(list)
        for n, к in карты:
            з = ключ(к)
            з = 'не указан' if з is None else з
            счёт[з] += 1
            номера[з].append(n)
        for з, сколько in счёт.most_common():
            ws.append([название, з, сколько, ', '.join(str(x) for x in номера[з])])
            строка += 1
            for i in range(1, 5):
                c = ws.cell(row=строка, column=i)
                c.alignment = Alignment(vertical='top', wrap_text=True)
                c.border = РАМКА
            ws.cell(row=строка, column=2).font = Font(bold=True)
        ws.append(['', f'разных значений: {len(счёт)}', sum(счёт.values()), ''])
        строка += 1
        for c in ws[строка]:
            c.font = Font(bold=True)
        ws.append([])
        строка += 1

    блок('что считает множитель', lambda к: что_считает(к.get('multiplier', '')))
    блок('ПО за единицу', lambda к: за_единицу(к.get('multiplier', '')))
    блок('потолок множителя', lambda к: потолок(к.get('multiplier', '')))
    ws.freeze_panes = 'A2'


def main():
    свод, в, набор, путь = версии.каталог('super_objectives')
    цель = f'docs/СУПЕР-ЗАДАНИЯ {в} — все карты.xlsx'
    карты = [(n, к) for n, к in enumerate(набор, start=1)]

    wb = Workbook()
    wb.remove(wb.active)
    лист_карт(wb, карты)
    сводка(wb, карты)
    wb.save(цель)
    print(f'готово: {цель}')
    print(f'  свод {свод}, каталог {в}, источник {путь}')
    print(f'  карт {len(карты)}, листов {len(wb.sheetnames)}')


if __name__ == '__main__':
    main()
