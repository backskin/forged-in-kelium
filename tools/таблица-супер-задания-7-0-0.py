# -*- coding: utf-8 -*-
"""ТАБЛИЦА СУПЕР-ЗАДАНИЙ 7.0.0 — семнадцать карт целиком.

Супер-задания 7.0 (решение дизайнера 04.09.2026) устроены иначе прежних:
карта НЕ СЖИГАЕТСЯ вовсе. Верх — множитель победных очков в финале, низ —
жёсткое требование с очень хорошей наградой. Игрок получает ОДНУ карту втайне
на раздаче (с 06.09.2026 тянется одна, правило «раздать 2, оставить 1»
отменено).

Пять карт (s5_13..s5_17) добавлены 04.09.2026, когда базовых источников очков
осталось шесть: монеты, несданные уничтоженные жетоны, здания и войска на поле очков больше
не дают, и всё это переехало сюда множителями. Ставки на картах щедрее прежних
базовых курсов нарочно — источник работает не у всех четверых, а только у того,
кто взял карту, и она у него одна на партию.

Данные берутся из data/cards/super_objectives.7.0.0.yaml — файл и есть источник.

Запуск: python tools/таблица-супер-задания-7-0-0.py
"""
import io

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ИСТОЧНИК = 'data/cards/super_objectives.7.0.0.yaml'
ЦЕЛЬ = 'docs/СУПЕР ЗАДАНИЯ 7.0.0 — карты целиком.xlsx'

ТОНКАЯ = Side(style='thin', color='FFB0B0B0')
РАМКА = Border(left=ТОНКАЯ, right=ТОНКАЯ, top=ТОНКАЯ, bottom=ТОНКАЯ)
# Карты 13-17 пришли из базовых источников очков — отмечены другим фоном,
# чтобы дизайнер видел новую пятёрку отдельно от старой дюжины.
СТАРЫЙ = 'FFF2EFF9'
НОВЫЙ = 'FFEFF6F2'
НОВЫЕ = {'s5_13', 's5_14', 's5_15', 's5_16', 's5_17'}


def main():
    данные = yaml.safe_load(io.open(ИСТОЧНИК, encoding='utf-8'))
    карты = данные['super_objectives']

    wb = Workbook()
    ws = wb.active
    ws.title = 'Супер задания 7.0.0'
    заголовки = ['№', 'id', 'НАЗВАНИЕ', 'ВЕРХ: множитель в финале',
                 'НИЗ: требование', 'НИЗ: награда', 'описание']
    ширины = [4, 9, 22, 46, 46, 46, 62]
    ws.append(заголовки)
    for i, (з, ш) in enumerate(zip(заголовки, ширины), start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor='FF3C3160')
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = РАМКА
        ws.column_dimensions[get_column_letter(i)].width = ш
    ws.row_dimensions[1].height = 34

    for n, к in enumerate(карты, start=1):
        ws.append([n, к['id'], к['name'], к.get('multiplier', ''),
                   к.get('requirement', ''), к.get('reward', ''),
                   к.get('описание', '')])
        фон = НОВЫЙ if к['id'] in НОВЫЕ else СТАРЫЙ
        for i in range(1, len(заголовки) + 1):
            c = ws.cell(row=n + 1, column=i)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = РАМКА
            c.fill = PatternFill('solid', fgColor=фон)
        ws.cell(row=n + 1, column=3).font = Font(bold=True)

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = f'A1:G{len(карты) + 1}'
    wb.save(ЦЕЛЬ)
    print(f'готово: {ЦЕЛЬ}, карт {len(карты)} '
          f'(из них новых по правке 04.09.2026: {len(НОВЫЕ)})')


if __name__ == '__main__':
    main()
