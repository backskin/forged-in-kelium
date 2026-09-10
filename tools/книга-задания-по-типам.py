# -*- coding: utf-8 -*-
"""КНИГА ЗАДАНИЙ — лист на каждое семейство плюс сводка эффектов.

Семейств семь (Выработка, Засада, Обеспечение, Развёртывание, Расплата,
Устранение, Экспансия), начальные задания идут отдельным листом. Разрез именно
такой, потому что решение «чего в наборе не хватает» принимается внутри
семейства, а не по алфавиту.

Сводный лист отвечает на вопрос «какой утиль где стоит»: сколько карт его
несут и НОМЕРА этих карт. Номера сквозные по всему каталогу и не сбиваются при
переходе с листа на лист.

Версия каталога НЕ ЗАШИТА: берётся из свода, который играет движок
(GameConfig.DEFAULT_RULESET -> content_versions.objectives), см. tools/версии.py.
Прежний генератор читал tools/таблица-задания-1-15-0.py и потому собирал книгу
из каталога 1.15.0, когда в игре давно стоял 1.17.0 — молча, без ошибок.

Семейства живут в классах kelium.cards.objectives и попадают в yaml
комментарием шапки; оттуда и читаются, второй копии списка нет.

Запуск: python tools/книга-задания-по-типам.py
"""
import sys
from collections import Counter, defaultdict

sys.path.insert(0, 'tools')
import версии                                    # noqa: E402
from openpyxl import Workbook                    # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side   # noqa: E402
from openpyxl.utils import get_column_letter     # noqa: E402

ТОНКАЯ = Side(style='thin', color='FFB0B0B0')
РАМКА = Border(left=ТОНКАЯ, right=ТОНКАЯ, top=ТОНКАЯ, bottom=ТОНКАЯ)

# Цвет листа на семейство — чтобы вкладки различались с одного взгляда.
ЦВЕТА = {
    'Выработка': 'FF2F5D50', 'Засада': 'FF6D2F2F', 'Обеспечение': 'FF2F4F6D',
    'Развёртывание': 'FF6D4C2F', 'Расплата': 'FF5A2F6D', 'Устранение': 'FF6D2F55',
    'Экспансия': 'FF3F5D2F', 'Начальные': 'FF3C3160',
}

РЕСУРСЫ = {'coin': 'монет', 'ammo': 'боеприпасов', 'kelium': 'келемия',
           'trophy': 'трофеев', 'objective_card': 'карт задания',
           'arsenal_card': 'карт арсенала', 'container': 'контейнеров',
           'module': 'жетон модуля', 'vp': 'ПО'}


def награда(d):
    """Словарь награды в человеческую строку: {'coin': 3} -> «3 монет»."""
    if not d:
        return ''
    куски = []
    for k, v in d.items():
        имя = РЕСУРСЫ.get(k, k)
        куски.append(f'{имя}: {v}' if isinstance(v, str) else f'{v} {имя}')
    return ' · '.join(куски)


def шапку(ws, заголовки, ширины, цвет):
    ws.append(заголовки)
    for i, (з, ш) in enumerate(zip(заголовки, ширины), start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color='FFFFFFFF')
        c.fill = PatternFill('solid', fgColor=цвет)
        c.alignment = Alignment(vertical='center', wrap_text=True)
        c.border = РАМКА
        ws.column_dimensions[get_column_letter(i)].width = ш
    ws.row_dimensions[1].height = 34


def лист(wb, имя, карты):
    цвет = ЦВЕТА.get(имя, 'FF2F5D50')
    ws = wb.create_sheet(имя[:31])
    ws.sheet_properties.tabColor = цвет[2:]
    заголовки = ['№', 'id', 'НАЗВАНИЕ', 'вид', 'УСЛОВИЕ', 'УСИЛЕНИЕ',
                 'базовая награда', 'сверхнаграда', 'УТИЛЬ (верх)', 'описание']
    шапку(ws, заголовки, [4, 7, 22, 10, 52, 38, 20, 20, 46, 58], цвет)
    for строка, (n, к) in enumerate(карты, start=2):
        верх = к.get('top') or {}
        ws.append([
            n, к['id'], к.get('name', ''), к.get('type', ''),
            (к.get('requirement') or {}).get('условие', ''),
            (к.get('enhanced') or {}).get('условие', ''),
            награда(к.get('base_reward')), награда(к.get('special_reward')),
            верх.get('label', ''), к.get('описание', ''),
        ])
        for i in range(1, len(заголовки) + 1):
            c = ws.cell(row=строка, column=i)
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.border = РАМКА
        ws.cell(row=строка, column=3).font = Font(bold=True)
    ws.append([])
    ws.append(['', f'карт: {len(карты)}'])
    ws.cell(row=len(карты) + 3, column=2).font = Font(bold=True)
    ws.freeze_panes = 'A2'


def сводка(wb, карты, семья):
    ws = wb.create_sheet('Эффекты — сводка')
    ws.sheet_properties.tabColor = '3C3160'
    шапку(ws, ['разрез', 'значение', 'карт', 'НОМЕРА КАРТ', 'пояснение'],
          [16, 40, 6, 34, 60], 'FF3C3160')
    строка = 1

    def блок(название, ключ_карты, подпись=None):
        nonlocal строка
        счёт, номера, поясн = Counter(), defaultdict(list), {}
        for n, к in карты:
            значение = ключ_карты(к)
            счёт[значение] += 1
            номера[значение].append(n)
            if подпись:
                поясн.setdefault(значение, подпись(к))
        for значение, сколько in счёт.most_common():
            ws.append([название, значение, сколько,
                       ', '.join(str(x) for x in номера[значение]),
                       поясн.get(значение, '')])
            строка += 1
            for i in range(1, 6):
                c = ws.cell(row=строка, column=i)
                c.alignment = Alignment(vertical='top', wrap_text=True)
                c.border = РАМКА
            ws.cell(row=строка, column=2).font = Font(bold=True)
        ws.append(['', f'разных значений: {len(счёт)}', sum(счёт.values()), '', ''])
        строка += 1
        for c in ws[строка]:
            c.font = Font(bold=True)
        ws.append([])
        строка += 1

    # Утиль — главный разрез: за ним и открывают эту книгу.
    блок('УТИЛЬ (верх)',
         lambda к: (к.get('top') or {}).get('label', '—').split(':')[0].strip(),
         lambda к: (к.get('top') or {}).get('label', ''))
    блок('эффект утиля', lambda к: (к.get('top') or {}).get('effect', '—'))
    блок('семейство', lambda к: семья.get(к['id'], 'начальные'))
    блок('вид', lambda к: к.get('type', '—'))
    блок('чем проверяется', lambda к: к.get('checked_by', '—'))
    ws.freeze_panes = 'A2'


def main():
    свод, в, набор, путь = версии.каталог('objectives')
    семья, порядок = версии.семейства(путь)
    цель = f'docs/ЗАДАНИЯ {в} — по семействам.xlsx'
    карты = [(n, к) for n, к in enumerate(набор, start=1)]

    неразмеченные = [к['id'] for _, к in карты
                     if к.get('kind') == 'regular' and к['id'] not in семья]
    if неразмеченные:
        raise SystemExit(f'нет семейства у карт: {неразмеченные}')

    wb = Workbook()
    wb.remove(wb.active)
    учтено = 0
    for имя in порядок:
        свои = [(n, к) for n, к in карты if семья.get(к['id']) == имя]
        лист(wb, имя, свои)
        учтено += len(свои)
    начальные = [(n, к) for n, к in карты if к.get('kind') == 'starting']
    лист(wb, 'Начальные', начальные)
    учтено += len(начальные)
    сводка(wb, карты, семья)

    if учтено != len(карты):
        raise SystemExit(f'по листам разошлось: {учтено} из {len(карты)} карт')
    wb.save(цель)
    print(f'готово: {цель}')
    print(f'  свод {свод}, каталог {в}, источник {путь}')
    print(f'  карт {len(карты)}, листов {len(wb.sheetnames)}')


if __name__ == '__main__':
    main()
